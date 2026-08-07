"""
金鹰集团数据管理中心 - 数据同步模块
解析集控中心xlsx文件，导入所有sheet到SQLite，计算公式列
"""
import openpyxl
import sqlite3
import os
import sys
from datetime import datetime, timedelta
from db import get_conn, init_db, reset_db, DB_PATH

# xlsx列号→数字
def ci(col_str):
    return openpyxl.utils.column_index_from_string(col_str)


#  sheet → DB表 映射配置 
# 每个sheet的列映射: (excel_col_letter, db_column_name, is_formula)
# is_formula=True 的列需要Python计算

SHEET_CONFIG = {
    '名称匹配': {
        'table': 'name_matching',
        'start_row': 2,
        'cols': [
            ('A', 'identity', False),
            ('B', 'pm_name', False),
            ('C', 'eng_name', False),
            ('D', 'eng_email', False),
            ('E', 'project_name', False),
            ('F', 'erp_code', False),
            ('G', 'project_123', False),
        ]
    },
    '身份匹配表': {
        'table': 'identity_matching',
        'start_row': 2,
        'cols': [
            ('A', 'identity', False),
            ('B', 'pm_name', False),
            ('C', 'eng_name', False),
            ('D', 'project_name', False),
            ('E', 'col_e', False),
            ('F', 'col_f', False),
            ('G', 'col_g', False),
            ('H', 'col_h', False),
        ]
    },
    '导出列表_归属项目': {
        'table': 'project_attribution',
        'start_row': 2,
        'cols': [('A', 'project_name', False), ('B', 'count', False)]
    },
    '导出列表_隐患状态_1': {
        'table': 'hazard_status_list',
        'start_row': 2,
        'cols': [('A', 'status_name', False), ('B', 'count', False)]
    },
    '商贸预算': {
        'table': 'commerce_budget',
        'start_row': 2,
        'cols': [
            ('A', 'subject', False), ('B', 'adjusted_budget', False),
            ('C', 'period', False), ('D', 'pending_cost', False),
            ('E', 'initial_budget', False), ('F', 'actual_cost', False),
            ('G', 'available_budget', False), ('H', 'revised_budget', False),
            ('I', 'execution_rate', False), ('J', 'unit', False),
            ('K', 'additional_budget', False),
        ]
    },
    '物业预算': {
        'table': 'property_budget',
        'start_row': 2,
        'cols': [
            ('A', 'subject', False), ('B', 'adjusted_budget', False),
            ('C', 'period', False), ('D', 'pending_cost', False),
            ('E', 'initial_budget', False), ('F', 'actual_cost', False),
            ('G', 'available_budget', False), ('H', 'revised_budget', False),
            ('I', 'execution_rate', False), ('J', 'unit', False),
            ('K', 'additional_budget', False),
        ]
    },
    '后付费登记表': {
        'table': 'postpaid_register',
        'start_row': 2,
        'cols': [
            ('A', 'meter_name', False), ('B', 'meter_no', False),
            ('C', 'room_id', False), ('D', 'remark', False),
            ('E', 'identity', False), ('F', 'project_name', False),  # F有缓存值
        ]
    },
    '已报备故障': {
        'table': 'reported_faults',
        'start_row': 2,
        'cols': [
            ('A', 'room_name', False), ('B', 'meter_name', False),
            ('C', 'balance', False), ('D', 'base_account', False),
            ('E', 'subsidy_account', False), ('F', 'identity', False),
            ('G', 'meter_type', False), ('H', 'reading', False),
            ('I', 'online_status', False), ('J', 'meter_no', False),
            ('K', 'branch', False), ('L', 'switch_detail', False),
            ('M', 'comm_time', False), ('N', 'gateway_no', False),
            ('O', 'channel_no', False), ('P', 'project_name', False),  # P有缓存值
        ]
    },
    '安全履职': {
        'table': 'safety_duty',
        'start_row': 2,
        'cols': [
            ('A', 'timestamp', False), ('B', 'plan_type', False),
            ('C', 'start_date', False), ('D', 'end_date', False),
            ('E', 'plan_status', False), ('F', 'plan_content', False),
            ('G', 'position', False), ('H', 'responsible', False),
            ('I', 'unit', False), ('J', 'completion', False),
            ('K', 'interval_days', False),
            ('L', 'countdown', True),  # =D2-NOW()
        ]
    },
    '隐患核查': {
        'table': 'hazard_inspection',
        'start_row': 2,
        'cols': [
            ('A', 'project_name', False),  # 有缓存值
            ('B', 'category', False), ('C', 'level', False),
            ('D', 'responsible', False), ('E', 'deadline', False),
            ('F', 'description', False), ('G', 'measure', False),
            ('H', 'audit_status', False), ('I', 'finish_time', False),
            ('J', 'hazard_status', False), ('K', 'delay_days', False),
            ('L', 'delay_note', False), ('M', 'creator', False),
            ('N', 'create_time', False),
            ('O', 'remaining_days', False),  # 有缓存值
            ('P', 'create_days', True),  # =NOW()-N2
        ]
    },
    '末端水压': {
        'table': 'water_pressure',
        'start_row': 2,
        'cols': [
            ('A', 'project_name', False),  # 有缓存值(ArrayFormula)
            ('B', 'device_no', False), ('C', 'net_signal', False),
            ('D', 'probe_battery', False), ('E', 'module_battery', False),
            ('F', 'pressure', False), ('G', 'temperature', False),
            ('H', 'report_time', False), ('I', 'device_name', False),
            ('J', 'install_loc', False), ('K', 'normal_status', False),
            ('L', 'online_status', False), ('M', 'col_m', False),
            ('N', 'col_n', False),
            ('O', 'comm_delay', True),  # =NOW()-H2
            ('P', 'pressure_fault', True),  # =IF(F<=0.02,"失压",IF(F>=1.3,"超压","0"))
        ]
    },
    'ruixin': {
        'table': 'ruixin_meters',
        'start_row': 2,
        'cols': [
            ('A', 'seq', False), ('B', 'region', False),
            ('C', 'building_id', False), ('D', 'building', False),
            ('E', 'floor_id', False), ('F', 'floor', False),
            ('G', 'device_id', False), ('H', 'timeout', False),
            ('I', 'room_id', False), ('J', 'room', False),
            ('K', 'unit', False), ('L', 'balance', False),
            ('M', 'base_balance', False), ('N', 'subsidy_balance', False),
            ('O', 'share_balance', False), ('P', 'elec_subsidy', False),
            ('Q', 'cold_water', False), ('R', 'hot_water', False),
            ('S', 'heating', False), ('T', 'identity', False),
            ('U', 'product_class', False), ('V', 'meter_type', False),
            ('W', 'comm_status', False), ('X', 'device_status', False),
            ('Y', 'install_addr', False), ('Z', 'switch_status', False),
            ('AA', 'switch_code', False), ('AB', 'channel_name', False),
            ('AC', 'channel', False), ('AD', 'channel_no', False),
            ('AE', 'display_val', False), ('AF', 'reading', False),
            ('AG', 'gateway_no', False), ('AH', 'product_id_tag', False),
            ('AI', 'product_id', False), ('AJ', 'comm_time', False),
            ('AK', 'data_time', False), ('AL', 'controllable', False),
            ('AM', 'writable', False), ('AN', 'atom_code', False),
            ('AO', 'control_mode', False), ('AP', 'pricing', False),
            ('AQ', 'balance_status', False), ('AR', 'status', False),
            ('AS', 'alarm_enabled', False), ('AT', 'alarm_threshold', False),
            ('AU', 'occupied', False), ('AV', 'location', False),
            ('AW', 'joint_control', False), ('AX', 'comm_type', False),
            ('AY', 'signal_strength', False), ('AZ', 'battery_status', False),
            ('BA', 'ct_ratio', False), ('BB', 'pt_ratio', False),
            ('BC', 'meter_no', False), ('BD', 'version', False),
            ('BE', 'biz_format', False),
            # BF 跳过(空列)
            ('BG', 'project_name', True),   # =VLOOKUP(T,名称匹配!A:E,5,0)
            ('BH', 'offline_hours', True),  # =NOW()-AK
            ('BI', 'is_postpaid', True),    # =COUNTIF(后付费登记表!C:C,I)>=1
            ('BJ', 'fault_tag', True),     # 复杂IF
            ('BK', 'fault_desc', True),    # 拼接
        ]
    },
    '123': {
        'table': 'equipment_123',
        'start_row': 2,
        'cols': [
            ('A', 'project_abbr', False), ('B', 'device_group', False),
            ('C', 'device_id', False), ('D', 'device_name', False),
            ('E', 'device_type', False), ('F', 'device_tag', False),
            ('G', 'online_status', False), ('H', 'device_status', False),
            ('I', 'sn_code', False), ('J', 'imei', False),
            ('K', 'mac', False), ('L', 'address', False),
            ('M', 'longitude', False), ('N', 'latitude', False),
            ('O', 'firmware', False), ('P', 'last_comm', False),
            ('Q', 'create_time', False), ('R', 'verify_code', False),
            ('S', 'protocol', False), ('T', 'account', False),
            ('U', 'update_time', False), ('V', 'battery', False),
            ('W', 'signal', False), ('X', 'humidity', False),
            ('Y', 'temperature', False), ('Z', 'ext_power', False),
            ('AA', 'report_time', False), ('AB', 'formaldehyde', False),
            ('AC', 'pm25', False), ('AD', 'co2', False),
            ('AE', 'tvoc', False), ('AF', 'water_immersion', False),
            ('AG', 'noise', False), ('AH', 'pressure', False),
            ('AI', 'project_name', True),  # =XLOOKUP(A,名称匹配!G:G,名称匹配!E:E)
            ('AJ', 'comm_delay', True),    # =NOW()-P
        ]
    },
}

# 同步顺序(依赖关系)
SYNC_ORDER = [
    '名称匹配', '身份匹配表',
    '导出列表_归属项目', '导出列表_隐患状态_1',
    '后付费登记表', '已报备故障',
    '商贸预算', '物业预算',
    '安全履职', '隐患核查', '末端水压',
    'ruixin', '123',
]


def _parse_number(val):
    """尝试转数字，失败返回原值"""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return val
    s = str(val).strip()
    if s == '' or s == 'None':
        return None
    # 处理百分比
    if '%' in s:
        try:
            return float(s.replace('%', '')) / 100
        except:
            return s
    try:
        return float(s)
    except:
        return s


def _parse_datetime(val):
    """解析日期时间"""
    if val is None:
        return None
    if isinstance(val, datetime):
        return val
    s = str(val).strip()
    if s == '' or s == 'None':
        return None
    for fmt in ['%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M:%S.%f', '%Y-%m-%d']:
        try:
            return datetime.strptime(s, fmt)
        except:
            pass
    return s  # 返回原字符串


def _excel_serial_to_dt(serial):
    """Excel序列号转datetime"""
    if isinstance(serial, (int, float)):
        try:
            return datetime(1899, 12, 30) + timedelta(days=serial)
        except:
            pass
    return None


def _get_col_val(ws_val, ws_f, row, col_idx):
    """从缓存值表获取值"""
    v = ws_val.cell(row=row, column=col_idx).value
    return v


def sync_sheet(ws_v, ws_f, config, conn, log_entries, source_file):
    """同步单个sheet到数据库"""
    table = config['table']
    start_row = config['start_row']
    cols = config['cols']
    max_row = ws_v.max_row

    # 获取DB列名(跳过公式列)
    db_cols = [c[1] for c in cols if not c[2]]
    formula_cols = [c for c in cols if c[2]]

    # 清空目标表
    conn.execute(f"DELETE FROM {table}")

    rows_data = []
    for row_idx in range(start_row, max_row + 1):
        # 跳过完全空的行
        row_has_data = False
        row = {}
        for col_letter, db_col, is_formula in cols:
            col_idx = ci(col_letter)
            if is_formula:
                row[db_col] = None  # 公式列后面计算
                continue
            val = _get_col_val(ws_v, ws_f, row_idx, col_idx)
            if val is not None:
                row_has_data = True
            row[db_col] = val
        if not row_has_data:
            continue
        rows_data.append(row)

    # 批量插入非公式列
    if db_cols and rows_data:
        placeholders = ', '.join(['?'] * len(db_cols))
        col_list = ', '.join(db_cols)
        batch = [tuple(r.get(c) for c in db_cols) for r in rows_data]
        conn.executemany(
            f"INSERT INTO {table} ({col_list}) VALUES ({placeholders})",
            batch
        )

    imported = len(rows_data)
    log_entries.append({
        'sync_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'source_file': source_file,
        'sheet_name': config.get('sheet_label', table),
        'rows_imported': imported,
        'status': 'success',
        'message': f'导入{imported}行'
    })
    return imported


def sync_budget_summary(ws_v, conn, source_file):
    """同步预算使用情况sheet(矩阵布局归一化为行记录)"""
    # 矩阵布局: row1=类别名(每4列一组), row2=子标题(当期用量/当期占比/当年用量/当年占比)
    # row3+=项目数据, col A=项目名称

    # 读取类别映射: (起始列号, 类别名)
    categories = []
    for c in range(2, ws_v.max_column + 1, 4):
        cat_name = ws_v.cell(row=1, column=c).value
        if cat_name and str(cat_name).strip() not in ('0', ''):
            categories.append((c, str(cat_name).strip()))

    conn.execute("DELETE FROM budget_summary")
    rows_data = []
    for r in range(3, ws_v.max_row + 1):
        proj = ws_v.cell(row=r, column=1).value
        if not proj or str(proj).strip() == '':
            continue
        proj = str(proj).strip()
        for start_col, cat_name in categories:
            cu = ws_v.cell(row=r, column=start_col).value      # 当期用量
            cr = ws_v.cell(row=r, column=start_col + 1).value  # 当期占比
            yu = ws_v.cell(row=r, column=start_col + 2).value  # 当年用量
            yr = ws_v.cell(row=r, column=start_col + 3).value  # 当年占比
            # 跳过全空的类别
            if cu is None and yu is None:
                continue
            rows_data.append((proj, cat_name, cu, str(cr) if cr is not None else None,
                              yu, str(yr) if yr is not None else None, r))

    if rows_data:
        conn.executemany(
            "INSERT INTO budget_summary (project_name, category, current_usage, current_rate, yearly_usage, yearly_rate, sort_order) VALUES (?,?,?,?,?,?,?)",
            rows_data
        )
    return len(rows_data)


def compute_formula_columns(conn):
    """计算所有公式列(在所有原始数据导入后执行)"""

    #  1. 构建 VLOOKUP 映射: identity → project_name 
    name_map_rows = conn.execute(
        "SELECT identity, project_name, project_123 FROM name_matching"
    ).fetchall()
    identity_to_project = {}
    project123_to_project = {}
    for r in name_map_rows:
        if r['identity']:
            identity_to_project[r['identity']] = r['project_name']
        if r['project_123']:
            project123_to_project[r['project_123']] = r['project_name']

    #  2. 构建后付费登记表 room_id 集合 
    postpaid_room_ids = set()
    for r in conn.execute("SELECT DISTINCT room_id FROM postpaid_register WHERE room_id IS NOT NULL"):
        postpaid_room_ids.add(str(r['room_id']).strip())

    #  3. 构建已报备故障 (project_name, meter_addr) 集合 
    reported_set = set()
    for r in conn.execute("SELECT DISTINCT project_name, meter_no FROM reported_faults"):
        if r['project_name'] and r['meter_no']:
            reported_set.add((r['project_name'].strip(), str(r['meter_no']).strip()))

    now = datetime.now()

    #  4. 计算 ruixin 公式列 
    print("  计算 ruixin 公式列...")
    ruixin_rows = conn.execute(
        "SELECT id, identity, room_id, data_time, location, switch_status, install_addr, room FROM ruixin_meters"
    ).fetchall()

    updates = []
    for r in ruixin_rows:
        rid = r['id']
        identity = r['identity'] or ''
        room_id = str(r['room_id']).strip() if r['room_id'] else ''
        data_time = r['data_time']
        location = r['location'] or ''
        switch_status = r['switch_status'] or ''
        install_addr = r['install_addr'] or ''
        room = r['room'] or ''

        # BG: 项目名称 = VLOOKUP(identity, 名称匹配!A:E, 5, 0)
        project_name = identity_to_project.get(identity, '')

        # BH: 离线时间 = NOW() - AK(data_time)
        offline_hours = None
        if data_time:
            dt = _parse_datetime(data_time)
            if dt and isinstance(dt, datetime):
                offline_hours = (now - dt).total_seconds() / 86400.0  # 天数

        # BI: 是否签呈后付 = COUNTIF(后付费登记表!C:C, room_id) >= 1
        is_postpaid = 1 if room_id in postpaid_room_ids else 0

        # BJ: 故障点 (复杂IF)
        fault_tag = ''
        # 先检查是否已报备
        is_reported = (project_name.strip(), install_addr.strip()) in reported_set if project_name and install_addr else False
        if not is_reported:
            # 【离线】: BH >= 1 (离线超过1天)
            if offline_hours is not None and offline_hours >= 1:
                fault_tag += '【离线】'
            # 【异常送电】: BI=0 AND 不含"公区" AND 不含"【不断电】" AND Z="手动开电"
            if (is_postpaid == 0
                    and '公区' not in location
                    and '【不断电】' not in identity
                    and switch_status == '手动开电'):
                fault_tag += '【异常送电】'
            # 【无签呈后付费】: 不含"【自动断电】" AND BI=0 AND 不含"公区"
            if ('【自动断电】' not in identity
                    and is_postpaid == 0
                    and '公区' not in location):
                fault_tag += '【无签呈后付费】'

        # BK: 故障描述
        if fault_tag:
            fault_desc = f'仪表号【{install_addr}】安装位置【{room}】仪表名称【{location}】故障描述{fault_tag}'
        else:
            fault_desc = '0'

        updates.append((project_name, offline_hours, is_postpaid, fault_tag, fault_desc, rid))

    conn.executemany(
        "UPDATE ruixin_meters SET project_name=?, offline_hours=?, is_postpaid=?, fault_tag=?, fault_desc=? WHERE id=?",
        updates
    )
    print(f"    ruixin: {len(updates)} 行公式列已计算")

    #  5. 计算 123 公式列 
    print("  计算 123 公式列...")
    equip_rows = conn.execute("SELECT id, project_abbr, last_comm FROM equipment_123").fetchall()
    updates = []
    for r in equip_rows:
        eid = r['id']
        abbr = r['project_abbr'] or ''
        last_comm = r['last_comm']

        # AI: 标准项目名称 = XLOOKUP(abbr, 名称匹配!G:G, 名称匹配!E:E)
        project_name = project123_to_project.get(abbr, abbr)

        # AJ: 通讯时差 = NOW() - P(last_comm)
        comm_delay = None
        if last_comm:
            dt = _parse_datetime(last_comm)
            if dt and isinstance(dt, datetime):
                comm_delay = (now - dt).total_seconds() / 86400.0

        updates.append((project_name, comm_delay, eid))

    conn.executemany(
        "UPDATE equipment_123 SET project_name=?, comm_delay=? WHERE id=?",
        updates
    )
    print(f"    123: {len(updates)} 行公式列已计算")

    #  6. 计算 安全履职 L: 结束倒计时 = D - NOW() 
    print("  计算 安全履职 公式列...")
    duty_rows = conn.execute("SELECT id, end_date FROM safety_duty").fetchall()
    updates = []
    for r in duty_rows:
        end_date = r['end_date']
        countdown = None
        if end_date:
            dt = _parse_datetime(end_date)
            if dt and isinstance(dt, datetime):
                countdown = (dt - now).total_seconds() / 86400.0
        updates.append((countdown, r['id']))
    conn.executemany("UPDATE safety_duty SET countdown=? WHERE id=?", updates)
    print(f"    安全履职: {len(updates)} 行公式列已计算")

    #  7. 计算 隐患核查 P: 创建天数 = NOW() - N 
    print("  计算 隐患核查 公式列...")
    hazard_rows = conn.execute("SELECT id, create_time FROM hazard_inspection").fetchall()
    updates = []
    for r in hazard_rows:
        ct = r['create_time']
        create_days = None
        if ct:
            dt = _parse_datetime(ct)
            if dt and isinstance(dt, datetime):
                create_days = (now - dt).total_seconds() / 86400.0
        updates.append((create_days, r['id']))
    conn.executemany("UPDATE hazard_inspection SET create_days=? WHERE id=?", updates)
    print(f"    隐患核查: {len(updates)} 行公式列已计算")

    #  8. 计算 末端水压 O, P 
    print("  计算 末端水压 公式列...")
    wp_rows = conn.execute("SELECT id, report_time, pressure FROM water_pressure").fetchall()
    updates = []
    for r in wp_rows:
        rt = r['report_time']
        pressure = r['pressure']

        # O: 通讯时间 = NOW() - H
        comm_delay = None
        if rt:
            dt = _parse_datetime(rt)
            if dt and isinstance(dt, datetime):
                comm_delay = (now - dt).total_seconds() / 86400.0

        # P: 压力故障 = IF(F<=0.02, "失压", IF(F>=1.3, "超压", "0"))
        if pressure is not None:
            try:
                p = float(pressure)
                if p <= 0.02:
                    pressure_fault = '失压'
                elif p >= 1.3:
                    pressure_fault = '超压'
                else:
                    pressure_fault = '0'
            except (ValueError, TypeError):
                pressure_fault = '0'
        else:
            pressure_fault = '0'

        updates.append((comm_delay, pressure_fault, r['id']))
    conn.executemany("UPDATE water_pressure SET comm_delay=?, pressure_fault=? WHERE id=?", updates)
    print(f"    末端水压: {len(updates)} 行公式列已计算")

    conn.commit()


def sync_from_xlsx(xlsx_path, do_reset=True, sheet_aliases=None):
    """从xlsx文件同步数据到数据库
    
    Args:
        xlsx_path: xlsx文件路径
        do_reset: 是否重置数据库(全量导入时True, 增量导入时False)
        sheet_aliases: sheet名别名映射 {实际sheet名: 期望sheet名}
                      例如 {'瑞信数据': 'ruixin', '设备信息汇总': '123'}
    """
    if not os.path.exists(xlsx_path):
        return {'success': False, 'error': f'文件不存在: {xlsx_path}'}

    print(f"\n{'='*50}")
    print(f"开始数据同步: {xlsx_path}")
    print(f"{'='*50}")

    init_db()
    if do_reset:
        reset_db()

    # 打开xlsx (缓存值 + 公式)
    wb_v = openpyxl.load_workbook(xlsx_path, data_only=True)
    wb_f = openpyxl.load_workbook(xlsx_path, data_only=False)

    conn = get_conn()
    log_entries = []
    total_rows = 0
    errors = []

    for sheet_name in SYNC_ORDER:
        if sheet_name not in SHEET_CONFIG:
            continue
        
        # 查找sheet: 先找原名, 再找别名映射
        actual_name = sheet_name
        if sheet_name not in wb_v.sheetnames:
            # 通过别名查找
            found = False
            if sheet_aliases:
                for actual, expected in sheet_aliases.items():
                    if expected == sheet_name and actual in wb_v.sheetnames:
                        actual_name = actual
                        found = True
                        break
            if not found:
                # 静默跳过缺失的sheet(增量导入时正常)
                continue

        config = dict(SHEET_CONFIG[sheet_name])
        config['sheet_label'] = sheet_name
        ws_v = wb_v[actual_name]
        ws_f = wb_f[actual_name]

        print(f"\n  同步 [{sheet_name}] → {config['table']}...")
        try:
            imported = sync_sheet(ws_v, ws_f, config, conn, log_entries, os.path.basename(xlsx_path))
            total_rows += imported
            print(f"    OK {imported} 行")
        except Exception as e:
            error_msg = f'{sheet_name}: {str(e)}'
            errors.append(error_msg)
            log_entries.append({
                'sync_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                'source_file': os.path.basename(xlsx_path),
                'sheet_name': sheet_name,
                'rows_imported': 0,
                'status': 'error',
                'message': str(e)
            })
            print(f"    ERR 错误: {e}")

    # 同步预算使用情况(矩阵布局)
    if '预算使用情况' in wb_v.sheetnames:
        print(f"\n  同步 [预算使用情况] → budget_summary...")
        try:
            imported = sync_budget_summary(wb_v['预算使用情况'], conn, os.path.basename(xlsx_path))
            total_rows += imported
            log_entries.append({
                'sync_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                'source_file': os.path.basename(xlsx_path),
                'sheet_name': '预算使用情况',
                'rows_imported': imported,
                'status': 'success',
                'message': f'导入{imported}行'
            })
            print(f"    OK {imported} 行")
        except Exception as e:
            errors.append(f'预算使用情况: {str(e)}')
            print(f"    ERR 错误: {e}")

    # 计算公式列
    print(f"\n  计算公式列...")
    try:
        compute_formula_columns(conn)
    except Exception as e:
        errors.append(f'公式计算: {str(e)}')
        print(f"    ERR 公式计算错误: {e}")
        import traceback
        traceback.print_exc()

    # 写入同步日志
    for entry in log_entries:
        conn.execute(
            "INSERT INTO sync_log (sync_time, source_file, sheet_name, rows_imported, status, message) VALUES (?, ?, ?, ?, ?, ?)",
            (entry['sync_time'], entry['source_file'], entry['sheet_name'],
             entry['rows_imported'], entry['status'], entry['message'])
        )

    conn.commit()
    conn.close()

    result = {
        'success': len(errors) == 0,
        'total_rows': total_rows,
        'sheets_synced': len(log_entries),
        'errors': errors,
        'sync_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'source_file': os.path.basename(xlsx_path),
    }
    print(f"\n{'='*50}")
    print(f"同步完成: {total_rows} 行, {len(log_entries)} 个sheet, {len(errors)} 个错误")
    print(f"{'='*50}\n")
    return result


def get_sync_status():
    """获取同步状态"""
    conn = get_conn()
    # 最新同步记录
    latest = conn.execute(
        "SELECT * FROM sync_log ORDER BY id DESC LIMIT 20"
    ).fetchall()

    # 各表行数
    tables = [
        'ruixin_meters', 'equipment_123', 'safety_duty', 'hazard_inspection',
        'water_pressure', 'commerce_budget', 'property_budget',
        'postpaid_register', 'reported_faults', 'name_matching',
        'identity_matching', 'project_attribution', 'hazard_status_list'
    ]
    table_counts = {}
    for t in tables:
        count = conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
        table_counts[t] = count

    last_sync = conn.execute(
        "SELECT sync_time FROM sync_log ORDER BY id DESC LIMIT 1"
    ).fetchone()

    conn.close()
    return {
        'last_sync': last_sync['sync_time'] if last_sync else None,
        'table_counts': table_counts,
        'recent_logs': [dict(r) for r in latest],
    }


if __name__ == '__main__':
    # 命令行测试
    xlsx_path = sys.argv[1] if len(sys.argv) > 1 else None
    if not xlsx_path:
        # 使用默认路径
        workspace = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
        xlsx_path = os.path.join(workspace, 'media', 'c4b8964a2e814b3bb99f5edf9f316ee4__总控表_集控表.xlsx')

    result = sync_from_xlsx(xlsx_path)
    print(f"\n同步结果: {result}")

    status = get_sync_status()
    print(f"\n同步状态:")
    print(f"  最后同步: {status['last_sync']}")
    for t, c in status['table_counts'].items():
        print(f"  {t}: {c} 行")
