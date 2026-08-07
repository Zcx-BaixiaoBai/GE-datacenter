"""
123物联网云平台设备信息爬取模块
系统: https://cloud.sd123iot.com
功能: 登录所有项目账号, 拉取全部设备信息
"""
import requests
import json
import hashlib
import time
import os
import urllib3
from datetime import datetime
from typing import List, Dict, Tuple

urllib3.disable_warnings()

BASE_URL = 'https://cloud.sd123iot.com:8443/usrCloud'

DATA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'data')
ACCOUNTS_FILE = os.path.join(DATA_DIR, 'sd123_accounts.xlsx')
CACHE_FILE = os.path.join(DATA_DIR, 'sd123_devices.json')


def _md5(text: str) -> str:
    return hashlib.md5(text.encode('utf-8')).hexdigest()


def load_accounts() -> List[Tuple[str, str, str]]:
    """读取账号表 -> [(项目名, 账号, 密码), ...]"""
    import openpyxl
    if not os.path.exists(ACCOUNTS_FILE):
        return []
    wb = openpyxl.load_workbook(ACCOUNTS_FILE, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    accounts = []
    for row in rows:
        if len(row) >= 3 and row[1] and row[2]:
            project = str(row[1]).strip()
            account = str(row[2]).strip()
            password = str(row[3]).strip() if len(row) > 3 else ''
            accounts.append((project, account, password))
    return accounts


def _login(account: str, password: str) -> str:
    """登录获取token"""
    session = requests.Session()
    session.verify = False
    r = session.post(f'{BASE_URL}/user/login',
                     json={'account': account, 'password': _md5(password)},
                     timeout=15)
    data = r.json()
    if data.get('status') == 0:
        return data['data']['token']
    return None


def _get_all_devices(token: str) -> List[Dict]:
    """获取所有设备列表"""
    headers = {'Content-Type': 'application/json', 'token': token}
    all_devices = []
    page_no = 1
    page_size = 500

    while True:
        body = {'pageNo': page_no, 'pageSize': page_size}
        r = requests.post(f'{BASE_URL}/dev/getDevs', headers=headers,
                          json=body, timeout=30)
        data = r.json()
        if data.get('status') != 0:
            break
        devs = data['data'].get('dev', [])
        all_devices.extend(devs)
        total = data['data'].get('total', 0)
        if len(all_devices) >= total or not devs:
            break
        page_no += 1

    return all_devices


def _process_device(dev: dict, project_name: str) -> dict:
    """提取设备关键字段"""
    online_status = '在线' if dev.get('onlineStatus') == 1 else '离线'

    device_status = dev.get('deviceStatus', {})
    status_parts = []
    if device_status.get('forbidden'):
        status_parts.append('禁用')
    if device_status.get('datapointAlarm'):
        status_parts.append('数据报警')
    if device_status.get('monitorAlarm'):
        status_parts.append('监控报警')
    if device_status.get('update'):
        status_parts.append('升级中')
    status_str = ','.join(status_parts) if status_parts else '正常'

    tags = dev.get('tags', [])
    tag_str = ','.join([t.get('tagName', '') for t in tags]) if tags else ''

    position = dev.get('position', '')
    lon, lat = '', ''
    if position and ',' in position:
        parts = position.split(',')
        lon, lat = parts[0], parts[1]

    online_time = ''
    if device_status.get('onlineTime'):
        try:
            online_time = datetime.fromtimestamp(
                device_status['onlineTime']
            ).strftime('%Y-%m-%d %H:%M:%S')
        except Exception:
            pass

    create_time = ''
    if dev.get('createTime'):
        try:
            create_time = datetime.fromtimestamp(
                dev['createTime']
            ).strftime('%Y-%m-%d %H:%M:%S')
        except Exception:
            pass

    return {
        '归属项目': project_name,
        '设备分组': dev.get('groupName', ''),
        '设备ID': dev.get('devid', ''),
        '设备名称': dev.get('name', ''),
        '设备类型': dev.get('templateName', ''),
        '设备标签': tag_str,
        '在线状态': online_status,
        '设备状态': status_str,
        'SN码': dev.get('sn', ''),
        'IMEI': dev.get('imei', ''),
        'MAC': dev.get('mac', ''),
        '地址': dev.get('address', ''),
        '经度': lon,
        '纬度': lat,
        '固件版本': device_status.get('version', ''),
        '最近通讯时间': online_time,
        '创建时间': create_time,
        '验证码': dev.get('verifyCode', ''),
        '协议': str(dev.get('protocol', '')),
        '设备账号': dev.get('account', ''),
    }


def crawl_all() -> Tuple[List[dict], int]:
    """爬取所有项目的全部设备, 返回 (设备列表, 项目数)"""
    accounts = load_accounts()
    if not accounts:
        raise FileNotFoundError(f"账号文件不存在: {ACCOUNTS_FILE}")

    all_devices = []
    projects_done = 0

    for project, account, password in accounts:
        token = _login(account, password)
        if not token:
            continue

        devs = _get_all_devices(token)
        for d in devs:
            all_devices.append(_process_device(d, project))
        projects_done += 1

    # 缓存到JSON
    os.makedirs(DATA_DIR, exist_ok=True)
    with open(CACHE_FILE, 'w', encoding='utf-8') as f:
        json.dump(all_devices, f, ensure_ascii=False, indent=2)

    return all_devices, projects_done


def load_cache() -> List[dict]:
    """读取缓存的设备数据"""
    if os.path.exists(CACHE_FILE):
        with open(CACHE_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    return []


def export_xlsx(devices: List[dict], filepath: str):
    """导出为xlsx"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    if not devices:
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "设备信息汇总"

    keys = []
    seen = set()
    for d in devices:
        for k in d.keys():
            if k not in seen:
                keys.append(k)
                seen.add(k)
    h_font = Font(bold=True, size=10, color="FFFFFF")
    h_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)

    for ci, k in enumerate(keys, 1):
        cell = ws.cell(row=1, column=ci, value=k)
        cell.font = h_font
        cell.fill = h_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for ri, d in enumerate(devices, 2):
        for ci, k in enumerate(keys, 1):
            cell = ws.cell(row=ri, column=ci, value=d.get(k, ''))
            cell.border = border

    for ci in range(1, len(keys) + 1):
        ws.column_dimensions[chr(64 + ci) if ci <= 26 else 'A'].width = 16

    ws.auto_filter.ref = f"A1:{chr(64 + min(len(keys), 26))}{len(devices)+1}"
    ws.freeze_panes = "A2"

    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    wb.save(filepath)
